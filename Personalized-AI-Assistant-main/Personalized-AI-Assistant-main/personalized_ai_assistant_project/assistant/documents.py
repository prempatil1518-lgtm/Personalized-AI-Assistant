from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from assistant.config import CHUNK_SIZE, CHUNK_OVERLAP

@dataclass
class TextDocument:
    text: str
    source: str

def read_file_bytes(name, data):
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return [TextDocument(page.extract_text() or "", f"{name} (page {i+1})")
                for i, page in enumerate(PdfReader(BytesIO(data)).pages)]
    if suffix == ".docx":
        doc = DocxDocument(BytesIO(data))
        return [TextDocument("\n".join(p.text for p in doc.paragraphs if p.text.strip()), name)]
    if suffix in {".txt", ".md"}:
        return [TextDocument(data.decode("utf-8", errors="replace"), name)]
    raise ValueError(f"Unsupported file type: {suffix}")

def load_uploaded_documents(files):
    docs = []
    for file in files:
        docs.extend(read_file_bytes(file.name, file.getvalue()))
    return [d for d in docs if d.text.strip()]

def load_folder(folder):
    docs = []
    for path in Path(folder).rglob("*"):
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx", ".txt", ".md"}:
            docs.extend(read_file_bytes(path.name, path.read_bytes()))
    return [d for d in docs if d.text.strip()]

def split_documents(documents, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks = []
    for doc in documents:
        text = " ".join(doc.text.split())
        start = 0
        while start < len(text):
            end = min(len(text), start + chunk_size)
            if text[start:end].strip():
                chunks.append(TextDocument(text[start:end].strip(), doc.source))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return chunks
